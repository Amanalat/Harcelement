# -*- coding: utf-8 -*-
from docx import Document
from docx.shared import Pt, Cm, Mm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

OUT = 'doc_textes_partenaires.docx'

C_GREEN  = RGBColor(0x2a,0x7a,0x50)
C_ORANGE = RGBColor(0xc0,0x90,0x30)
C_RED    = RGBColor(0xc0,0x40,0x30)
C_GREY   = RGBColor(0x66,0x66,0x66)
C_DIM    = RGBColor(0x99,0x99,0x99)
C_LIGHT  = RGBColor(0xaa,0xaa,0xaa)
C_BLUE   = RGBColor(0x2a,0x4a,0x7a)
SPK = {
    'leo':RGBColor(0x1a,0x3a,0x6a),'ines':RGBColor(0x4a,0x1a,0x6a),
    'lucas':RGBColor(0x1a,0x5a,0x28),'jade':RGBColor(0x7a,0x3a,0x00),
    'theo':RGBColor(0x2a,0x4a,0x7a),'marine':RGBColor(0x7a,0x1a,0x4a),
    'kevin':RGBColor(0x55,0x55,0x55),'tante':RGBColor(0x5a,0x35,0x10),
    'clara':RGBColor(0x6a,0x1a,0x2a),'narr':RGBColor(0xaa,0xaa,0xaa),
    'lea':RGBColor(0x44,0x88,0x44),
}
CHOICE_BG  = {'good':'F4FAF7','ok':'FDFAF0','bad':'FDF4F4','neutral':'FAFAFA'}
CHOICE_COL = {'good':C_GREEN,'ok':C_ORANGE,'bad':C_RED,'neutral':C_GREY}
CHOICE_HEX = {'good':'2a7a50','ok':'c09030','bad':'c04030','neutral':'555555'}

def _pPr(p): return p._p.get_or_add_pPr()
def set_spacing(para,before=0,after=80,line=276):
    sp=OxmlElement('w:spacing')
    sp.set(qn('w:before'),str(int(before)))
    sp.set(qn('w:after'),str(int(after)))
    sp.set(qn('w:line'),str(int(line)))
    sp.set(qn('w:lineRule'),'auto')
    _pPr(para).append(sp)
def set_indent(para,left=0,hanging=0):
    ind=OxmlElement('w:ind')
    if left: ind.set(qn('w:left'),str(int(left)))
    if hanging: ind.set(qn('w:hanging'),str(int(hanging)))
    _pPr(para).append(ind)
def left_border(para,color_hex,sz=18):
    pBdr=OxmlElement('w:pBdr'); left=OxmlElement('w:left')
    left.set(qn('w:val'),'single'); left.set(qn('w:sz'),str(sz))
    left.set(qn('w:space'),'4'); left.set(qn('w:color'),color_hex)
    pBdr.append(left); _pPr(para).append(pBdr)
def bottom_border(para,color_hex='cccccc',sz=4):
    pBdr=OxmlElement('w:pBdr'); b=OxmlElement('w:bottom')
    b.set(qn('w:val'),'single'); b.set(qn('w:sz'),str(sz))
    b.set(qn('w:space'),'1'); b.set(qn('w:color'),color_hex)
    pBdr.append(b); _pPr(para).append(pBdr)
def shading(para,fill_hex):
    shd=OxmlElement('w:shd'); shd.set(qn('w:val'),'clear')
    shd.set(qn('w:color'),'auto'); shd.set(qn('w:fill'),fill_hex)
    _pPr(para).append(shd)
def keep_together(para): _pPr(para).append(OxmlElement('w:keepLines'))
def page_break_before(para): _pPr(para).append(OxmlElement('w:pageBreakBefore'))
def add_run(para,text,bold=False,italic=False,color=None,size=None,font=None):
    r=para.add_run(text)
    if bold: r.bold=True
    if italic: r.italic=True
    if color: r.font.color.rgb=color
    if size: r.font.size=Pt(size)
    if font: r.font.name=font
    return r

doc=Document()
for sec in doc.sections:
    sec.page_width=Mm(210); sec.page_height=Mm(297)
    sec.left_margin=Cm(2.5); sec.right_margin=Cm(2.5)
    sec.top_margin=Cm(2.2); sec.bottom_margin=Cm(2.2)
for pp in doc.paragraphs: pp._element.getparent().remove(pp._element)

def p(text='',align=WD_ALIGN_PARAGRAPH.LEFT,before=0,after=80,size=11,
      bold=False,italic=False,color=None,font=None,indent=0,line=276):
    para=doc.add_paragraph(); para.alignment=align
    set_spacing(para,before,after,line)
    if indent: set_indent(para,left=indent)
    if text: add_run(para,text,bold=bold,italic=italic,color=color,size=size,font=font)
    return para

def heading2(text,before=160,after=80):
    para=doc.add_paragraph(); set_spacing(para,before=before,after=after)
    add_run(para,text.upper(),bold=True,size=10,color=RGBColor(0x44,0x44,0x44),font='Arial')
    return para

def heading3(text,before=100,after=60):
    para=doc.add_paragraph(); set_spacing(para,before=before,after=after)
    add_run(para,text,bold=True,size=10,color=C_GREY,font='Arial')
    return para

def part_header(num_text,title_text):
    np=doc.add_paragraph(); page_break_before(np)
    set_spacing(np,before=0,after=40); shading(np,'1a1a1a')
    add_run(np,num_text.upper(),bold=False,size=8,color=RGBColor(0x88,0x88,0x88),font='Arial')
    tp=doc.add_paragraph(); set_spacing(tp,before=0,after=160); shading(tp,'1a1a1a')
    add_run(tp,title_text,bold=False,italic=True,size=17,color=RGBColor(0xf0,0xec,0xe4))

def phase_sep(text):
    para=doc.add_paragraph(); set_spacing(para,before=120,after=80)
    add_run(para,f'- {text} -',italic=True,size=8,color=C_DIM,font='Arial')
    para.alignment=WD_ALIGN_PARAGRAPH.CENTER

def time_sep(text):
    para=doc.add_paragraph(); set_spacing(para,before=80,after=60)
    add_run(para,text,italic=True,size=8.5,color=C_DIM,font='Arial')
    para.alignment=WD_ALIGN_PARAGRAPH.CENTER

def dialogue_line(spk_key,spk_label,text,italic_text=False):
    para=doc.add_paragraph(); set_spacing(para,before=0,after=30,line=276)
    keep_together(para)
    add_run(para,f'{spk_label:<12}',bold=True,size=8.5,
            color=SPK.get(spk_key,RGBColor(0x1a,0x1a,0x1a)),font='Arial Narrow')
    para.add_run('  ').font.size=Pt(8)
    add_run(para,text,italic=italic_text,size=10.5,
            color=C_GREY if italic_text else None)

def choice_block(kind,letter,text,reply=None,indent_left=700):
    para=doc.add_paragraph(); set_spacing(para,before=0,after=30,line=276)
    set_indent(para,left=indent_left); left_border(para,CHOICE_HEX[kind],sz=18)
    shading(para,CHOICE_BG[kind]); keep_together(para)
    if letter: add_run(para,f'{letter}  ',bold=True,size=8.5,color=CHOICE_COL[kind],font='Arial')
    add_run(para,text,size=10.5)
    if reply:
        para.add_run('\n')
        add_run(para,reply,italic=True,size=9,color=C_GREY)

def ending_block(kind,title,*lines):
    tp=doc.add_paragraph(); set_spacing(tp,before=80,after=20)
    set_indent(tp,left=300); left_border(tp,CHOICE_HEX.get(kind,'555555'),sz=24)
    shading(tp,CHOICE_BG.get(kind,'FAFAFA'))
    add_run(tp,title,bold=True,size=10.5,color=CHOICE_COL.get(kind,C_GREY))
    for line in lines:
        lp=doc.add_paragraph(); set_spacing(lp,before=0,after=20)
        set_indent(lp,left=300); left_border(lp,CHOICE_HEX.get(kind,'555555'),sz=24)
        shading(lp,CHOICE_BG.get(kind,'FAFAFA'))
        if line.startswith('«') or line.startswith('"') or line.startswith('-'):
            add_run(lp,line,italic=True,size=10,color=C_GREY)
        else: add_run(lp,line,size=10.5)

def quiz_block(num_text,question,options,correct_idx,explanation):
    tp=doc.add_paragraph(); set_spacing(tp,before=120,after=30)
    set_indent(tp,left=200); shading(tp,'F5F7FA'); left_border(tp,'2a4a7a',sz=8)
    add_run(tp,num_text,bold=True,size=10,color=C_BLUE,font='Arial')
    qp=doc.add_paragraph(); set_spacing(qp,before=0,after=60)
    set_indent(qp,left=200); shading(qp,'F5F7FA'); left_border(qp,'2a4a7a',sz=8)
    add_run(qp,question,size=10.5)
    for i,opt in enumerate(options):
        op=doc.add_paragraph(); set_spacing(op,before=0,after=20)
        set_indent(op,left=600); shading(op,'F5F7FA')
        if i==correct_idx: add_run(op,opt,bold=True,size=10,color=C_GREEN)
        else: add_run(op,opt,size=10)
    ep=doc.add_paragraph(); set_spacing(ep,before=40,after=80)
    set_indent(ep,left=200); shading(ep,'F5F7FA'); left_border(ep,'2a4a7a',sz=8)
    add_run(ep,explanation,italic=True,size=9.5,color=C_GREY)

def recovery_block(aunt_line,good_txt,good_reply,bad_txt,bad_reply):
    hp=doc.add_paragraph(); set_spacing(hp,before=60,after=20)
    set_indent(hp,left=600); left_border(hp,'e0c080',sz=12); shading(hp,'FFFBF0')
    add_run(hp,'Rattrapage possible (si mauvaise réponse) :',italic=True,size=9,
            color=RGBColor(0xa0,0x80,0x30),font='Arial')
    ap=doc.add_paragraph(); set_spacing(ap,before=0,after=20)
    set_indent(ap,left=600); left_border(ap,'e0c080',sz=12); shading(ap,'FFFBF0')
    add_run(ap,'La tante   ',bold=True,size=8.5,color=SPK['tante'],font='Arial Narrow')
    add_run(ap,aunt_line,size=10)
    gp=doc.add_paragraph(); set_spacing(gp,before=0,after=20)
    set_indent(gp,left=600); left_border(gp,CHOICE_HEX['good'],sz=14); shading(gp,CHOICE_BG['good'])
    add_run(gp,'ok  ',bold=True,size=9,color=C_GREEN,font='Arial')
    add_run(gp,good_txt,size=10); gp.add_run('\n')
    add_run(gp,good_reply,italic=True,size=9,color=C_GREY)
    bp=doc.add_paragraph(); set_spacing(bp,before=0,after=80)
    set_indent(bp,left=600); left_border(bp,CHOICE_HEX['bad'],sz=14); shading(bp,CHOICE_BG['bad'])
    add_run(bp,'x   ',bold=True,size=9,color=C_RED,font='Arial')
    add_run(bp,bad_txt,size=10); bp.add_run('\n')
    add_run(bp,bad_reply,italic=True,size=9,color=C_GREY)

def footnote_line(text):
    para=doc.add_paragraph(); set_spacing(para,before=200,after=0)
    para.alignment=WD_ALIGN_PARAGRAPH.CENTER; bottom_border(para,'eeeeee',sz=4)
    add_run(para,text,size=8,color=C_LIGHT,font='Arial')

def bullet(text,indent=300):
    para=doc.add_paragraph(); set_spacing(para,before=0,after=30)
    set_indent(para,left=indent)
    add_run(para,'*  ',bold=True,size=10,color=C_DIM)
    add_run(para,text,size=10.5)

def dm_block(account,time_label,harcel_type,messages):
    header=doc.add_paragraph(); set_spacing(header,before=60,after=10)
    set_indent(header,left=200); left_border(header,'3a4a6a',sz=14); shading(header,'F0F4FA')
    add_run(header,account,bold=True,size=9.5,color=C_BLUE,font='Arial')
    add_run(header,f'  {time_label}',size=8,color=C_DIM,font='Arial')
    if harcel_type:
        add_run(header,f'   [{harcel_type}]',italic=True,size=8,color=C_ORANGE,font='Arial')
    for msg in messages:
        mp=doc.add_paragraph(); set_spacing(mp,before=0,after=8,line=260)
        set_indent(mp,left=400); shading(mp,'F7F9FD')
        add_run(mp,f'"{msg}"',size=9.5)

def lea_line(side,text):
    para=doc.add_paragraph(); set_spacing(para,before=0,after=20,line=276)
    keep_together(para)
    lbl='Léa         ' if side=='lea' else 'Clara       '
    col=SPK['lea'] if side=='lea' else SPK['clara']
    add_run(para,lbl,bold=True,size=8.5,color=col,font='Arial Narrow')
    para.add_run('  ').font.size=Pt(8)
    add_run(para,text,size=10.5)

